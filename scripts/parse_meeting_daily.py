"""从指标日报 Excel 抽分炉日均值（入炉加权），覆盖 9/1 温压氧排烟常数。

行 20-22 为三炉：入炉(D) 主汽温(G) 主汽压(J) 给水(M) 蒸发量(P) 排烟(V) 氧量(Y)。
停炉（压力<1 或温度<200）不参与加权。
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

RAW = Path(__file__).resolve().parents[1] / "data" / "snmis_raw"
OUT = Path(__file__).resolve().parents[1] / "data" / "plant"

# 1-based excel cols
COL = dict(mass=4, steam_t=7, steam_p=10, feedwater_t=13, evap=16, fluegas_t=22, o2=25)
BOILER_ROWS = (20, 21, 22)


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_file(path: Path) -> dict | None:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    recs = []
    for r in BOILER_ROWS:
        mass = _num(ws.cell(r, COL["mass"]).value) or 0.0
        p = _num(ws.cell(r, COL["steam_p"]).value)
        t = _num(ws.cell(r, COL["steam_t"]).value)
        if p is None or t is None or p < 1.0 or t < 200:
            continue
        w = mass if mass > 0 else (_num(ws.cell(r, COL["evap"]).value) or 0.0)
        if w <= 0:
            w = 1.0
        recs.append({
            "w": w,
            "steam_p": p,
            "steam_t": t,
            "feedwater_t": _num(ws.cell(r, COL["feedwater_t"]).value),
            "o2": _num(ws.cell(r, COL["o2"]).value),
            "fluegas_t": _num(ws.cell(r, COL["fluegas_t"]).value),
        })
    if not recs:
        return None
    tw = sum(x["w"] for x in recs)
    out = {"n_on": len(recs)}
    for k in ("steam_p", "steam_t", "feedwater_t", "o2", "fluegas_t"):
        nums = [(x["w"], x[k]) for x in recs if x[k] is not None]
        out[k] = (sum(w * v for w, v in nums) / sum(w for w, _ in nums)) if nums else None
    out["w_sum"] = tw
    return out


def collect(glob: str) -> pd.DataFrame:
    rows = []
    for p in sorted(RAW.glob(glob)):
        # meet_p1_2026-09-01.xlsx
        iso = p.stem.split("_", 2)[-1]
        rec = parse_file(p)
        if rec is None:
            continue
        rec["time"] = iso
        rows.append(rec)
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df["time"] = pd.to_datetime(df["time"])
    return df.sort_values("time")


def main() -> int:
    OUT.mkdir(parents=True, exist_ok=True)
    cols = ["time", "steam_p", "steam_t", "feedwater_t", "o2", "fluegas_t", "n_on"]
    p1 = collect("meet_p1_2026-*.xlsx")
    p2 = collect("meet_p2_2026-*.xlsx")
    h2 = collect("meet_p2_2024-*.xlsx")
    if not p1.empty:
        p1[cols].to_csv(OUT / "meeting_p1_daily.csv", index=False, encoding="utf-8-sig")
    if not p2.empty:
        p2[cols].to_csv(OUT / "meeting_p2_daily.csv", index=False, encoding="utf-8-sig")
    if not h2.empty:
        h2[cols].to_csv(OUT / "meeting_p2_hist_daily.csv", index=False, encoding="utf-8-sig")
    print("p1", len(p1), "p2", len(p2), "hist", len(h2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
