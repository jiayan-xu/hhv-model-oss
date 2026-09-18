# -*- coding: utf-8 -*-
"""V1.1 → V1.2：测算结果页口径；PFAiX 1.0.64。"""
from pathlib import Path

from docx import Document

SRC = Path(__file__).resolve().parents[1] / "docs" / "twin_plan_v1.1.docx"
OUT = Path(__file__).resolve().parents[1] / "docs" / "twin_plan_v1.2.docx"
MAIN = Path(__file__).resolve().parents[1] / "docs" / "twin_plan.docx"

d = Document(str(SRC))


def set_para_text(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def replace_in_para(p, old, new):
    if old not in p.text:
        return False
    set_para_text(p, p.text.replace(old, new))
    return True


def set_cell(table, r, c, text):
    cell = table.rows[r].cells[c]
    if cell.paragraphs:
        set_para_text(cell.paragraphs[0], text)
        for extra in cell.paragraphs[1:]:
            set_para_text(extra, "")
    else:
        cell.text = text


# 全文替换版本号（封面+页眉类单元格）
for p in d.paragraphs:
    replace_in_para(p, "文档版本：V1.1", "文档版本：V1.2")
for tb in d.tables:
    for row in tb.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                replace_in_para(p, "文档版本：V1.1", "文档版本：V1.2")

# 摘要：补展示层
for p in d.paragraphs:
    replace_in_para(
        p,
        "本报告同时申请 DCS 历史数据只读访问权限——这是全链路中唯一需要厂内配合的前置条件。",
        "智浦助手数据模型页（1.0.64）按测算报告展示：接入了什么数据、采用什么算法、算出什么数；"
        "主口径后台固定，用户不选配置文件、不看到 yaml 文件名。"
        "本报告同时申请 DCS 历史数据只读访问权限——这是全链路中唯一需要厂内配合的前置条件。",
    )
    replace_in_para(p, "三、L1 热值软测量", "三、L1 热值测算小时化")
    replace_in_para(
        p,
        "2026-09-07 修订：DCS 小时接入反推、软测量 ML 停用。",
        "2026-09-07 修订：DCS 小时接入反推、软测量 ML 停用；"
        "智浦助手 1.0.64 测算结果页（不暴露配置文件）。",
    )

# 表 1 资产：PFAiX 行
set_cell(
    d.tables[1],
    6,
    0,
    "PFAiX 数据模型页",
)
set_cell(d.tables[1], 6, 1, "已上线 1.0.64")
set_cell(
    d.tables[1],
    6,
    2,
    "页签：测算结果 / 验证 / 定价 / 数据接入。"
    "测算结果展示接入数据与算法（蒸汽正平衡、小时筛选、交叉验证闸门、稳健分解），"
    "主数字为综合固废入炉热值；后台固定主口径，不暴露配置文件名，无演示数据/输出目录。软测量页签已撤。",
)

# 表 2 L1 名称
set_cell(d.tables[2], 2, 0, "L1 热值测算小时化")
set_cell(d.tables[2], 2, 1, "物理反推小时化；ML 学厂填停用")

# 表 9 本周里程碑：补 UI
set_cell(
    d.tables[9],
    1,
    1,
    "DCS 小时库接入反推主链；学厂填软测量停用；测算结果页上线（1.0.64）",
)

d.save(str(OUT))
print("saved", OUT)

# 尽量盖回正式稿（WPS 占用则跳过）
try:
    d.save(str(MAIN))
    print("overwrote", MAIN)
except OSError as e:
    print("MAIN locked, skip:", e)
