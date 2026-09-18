# -*- coding: utf-8 -*-
"""提交版（中和）：约 4 页 Word + 8 页 PPT。比短版有算法和验证，比长版少周明细。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt as PPt

OUT = Path(__file__).resolve().parents[1] / "outputs" / "config.snmis_daily"
DOCX = OUT / "固废热值测算说明_提交版.docx"
PPTX = OUT / "固废热值测算说明_提交版.pptx"

import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from hhv.facts import load_facts  # noqa: E402

# 结论数字一律从 outputs/config.snmis_daily/result.json 读取（缺文件即报错，不出旧数）
F = load_facts(OUT)

# 以下数字暂无可机读来源（来自入厂台账/化验台账的人工推导），集中在此处便于随数据更新核对；
# 每项标注来源，更新数据后请人工复核。
NARRATIVE = {
    "whitelist_t": 80300,        # 窗口内白名单固废（吨）— 入厂台账人工汇总
    "basket_t": 66797,           # 剔克劳丽后用于反推装修的篮子（吨）— 同上
    "kelali_days": 18,           # 克劳丽剔除天数
    "kelali_cars": 387,          # 克劳丽剔除车次
    "kelali_t": 8124,            # 克劳丽剔除吨数
    "with_kelali_kcal": 1309,    # 含克劳丽时综合数（对照，大卡）
    "decor_point_kj": 3375,      # 装修反推点估计（kJ/kg）
    "decor_point_kcal": 806,     # 装修反推点估计（大卡）
    "decor_c_kcal": 733,         # 解法 C 对应装修（大卡）
    "decor_pin2_kcal": 830,      # 只钉造纸+农林时的装修（大卡）
    "decor_range_kcal": "540～1100",  # 装修区间（大卡）
    "lab_decor_kcal": "4438～5747",   # 装修化验（筛上干样，大卡）
    "mismatch_weeks": 7,         # 物料对不上的周数
    "mismatch_move_kj": 182,     # 丢掉这些周后点估计移动（kJ/kg）
    "decor_t": 40949,            # 装修入炉吨数
}

FONT = "微软雅黑"
INK = RGBColor(0x2F, 0x34, 0x37)
MUTED = RGBColor(0x78, 0x77, 0x74)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
C_INK = PptRGB(0x2F, 0x34, 0x37)
C_MUTE = PptRGB(0x78, 0x77, 0x74)
C_WHITE = PptRGB(0xFF, 0xFF, 0xFF)
C_DARK = PptRGB(0x1C, 0x1F, 0x21)
C_BG = PptRGB(0xF7, 0xF6, 0xF3)
C_LINE = PptRGB(0xEA, 0xEA, 0xEA)
C_GREEN = PptRGB(0x34, 0x65, 0x38)
W, H = Emu(12192000), Emu(6858000)


def font(run, size=11, bold=False, color=INK, name=FONT):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rPr.get_or_add_rFonts().set(qn("w:eastAsia"), name)
    rPr.get_or_add_rFonts().set(qn("w:ascii"), name)
    rPr.get_or_add_rFonts().set(qn("w:hAnsi"), name)


def p(doc, text, size=11, bold=False, space=6, color=INK, before=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space)
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.line_spacing = 1.28
    run = para.add_run(text)
    font(run, size=size, bold=bold, color=color)
    return para


def shade(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shd)


def borders(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "C8C8C8")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def table(doc, header, rows, highlight_last=False):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.autofit = True
    for i, h in enumerate(header):
        cell = t.cell(0, i)
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        font(run, size=9, bold=True, color=WHITE)
        shade(cell, "1C1F21")
        borders(cell)
    for ri, row in enumerate(rows):
        last = highlight_last and ri == len(rows) - 1
        for ci, val in enumerate(row):
            cell = t.cell(ri + 1, ci)
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            font(run, size=9, bold=last)
            if last:
                shade(cell, "EDF3EC")
            borders(cell)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(8)


def build_docx():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.1)
    sec.top_margin = sec.bottom_margin = Cm(1.7)

    p(doc, "一般固废入炉热值测算说明", size=18, bold=True, space=2)
    p(doc, f"提交版  ·  {F['generated']}  ·  窗口 2026-02-01～09-01  ·  不含克劳丽", size=10, color=MUTED, space=12)

    p(doc, "一、结论", size=14, bold=True, space=8, before=4)
    table(doc, ["项目", "热值", "说明"], [
        ["综合固废（主报价）", f"{F['kcal_b']} 大卡/公斤（{F['q_b']:.0f} kJ/kg）",
         f"解法 B；区间约 {F['q_b_lo']/4.186:.0f}～{F['q_b_hi']/4.186:.0f} 大卡"],
        ["装修入炉（领导口径）", f"约 {NARRATIVE['decor_point_kcal']} 大卡/公斤（约 {NARRATIVE['decor_point_kj']} kJ/kg）",
         f"点估计 {NARRATIVE['decor_point_kcal']} 大卡；区间约 {NARRATIVE['decor_range_kcal']} 大卡"],
        ["同期生活垃圾", f"约 {F['kcal_msw']} 大卡/公斤", f"一期尺子 × 刻度 k={F['k']:.2f}"],
    ], highlight_last=False)
    p(doc, f"{F['kcal_b']} 大卡谈「这一篮子固废值多少热」。{NARRATIVE['decor_point_kcal']} 大卡单独看装修：造纸（其他）、农林按检测报告钉住，剩下的热给装修。两个数不要混。", size=11, space=10)

    p(doc, "二、数据与口径", size=14, bold=True, space=8, before=6)
    p(doc, f"工业固废取运营台白名单（货名含「其他」，跳过环卫所、城东）。大件粉碎、生活源、厨余、机扫算生活垃圾，不进本口径。窗口内白名单固废 {NARRATIVE['whitelist_t']/10000:.2f} 万 t；剔克劳丽日后用于反推装修的篮子为 {NARRATIVE['basket_t']/10000:.2f} 万 t。", size=11, space=6)
    p(doc, "热量：生产日报一给蒸汽流量，指标日报给温度、压力、氧量、排烟温度。只读导出，未改 SNMIS 填报。历史尺子用 2024 年二期纯烧，不用 2025（当年已掺烧）。", size=11, space=6)
    p(doc, f"克劳丽 {NARRATIVE['kelali_days']} 天、{NARRATIVE['kelali_cars']} 车、{NARRATIVE['kelali_t']} 吨整日剔除（杠杆大、无自家氧弹）。同一方法含克劳丽时综合数为 {NARRATIVE['with_kelali_kcal']} 大卡，仅作对照。", size=11, space=10)

    p(doc, "三、综合数怎么算", size=14, bold=True, space=8, before=6)
    p(doc, "锅炉烧掉的是混料。蒸汽热量和入炉吨数能算出「这一锅有多热」；算不出的是固废自己有多热。一期几乎不掺工业固废，当生活垃圾尺子。固废掺得越多、混料偏离尺子越多，把固废拧出来。", size=11, space=6)
    p(doc, "混料热值 = 生活垃圾热值 × (1 − 掺烧比例) + 固废热值 × 掺烧比例。", size=11, bold=True, space=6)
    p(doc, f"步骤：蒸汽热 ÷ 锅炉效率（{min(F['eta_p1'], F['eta_p2'])*100:.0f}%～{max(F['eta_p1'], F['eta_p2'])*100:.0f}%）÷ 入炉吨 → 日混料热值；筛掉低负荷、缺测点、克劳丽日，一周有效日不满 5 天不进回归；按周算掺烧比例（{F['n_weeks']} 周从 {F['alpha_min_pct']:.0f}% 到 {F['alpha_max_pct']:.0f}%，均 {F['alpha_mean_pct']:.0f}%）；一期当周乘 k={F['k']:.2f} 换到二期表计；稳健回归拧出固废热值。", size=11, space=8)
    table(doc, ["解法", "综合固废", "角色"], [
        ["B  一期同周 × k", f"{F['q_b']:.0f} kJ / {F['kcal_b']} 大卡", "主报"],
        ["C  2024 同炉周对齐", f"{F['q_c']:.0f} kJ / {F['kcal_c']} 大卡",
         f"并报，与 B 差 {abs(F['kcal_b']-F['kcal_c'])} 大卡"],
        ["A  自由回归", f"{F['q_a']:.0f} kJ / {F['kcal_a']} 大卡", "校验，区间太宽，不单独报"],
    ])
    p(doc, f"内部验证通过：B 与 C 相差 {F['b_minus_c']:.0f} kJ（门限 500）；高掺烧周隐含值仍落在区间内；丢掉物料对不上的 {NARRATIVE['mismatch_weeks']} 周后再算，只挪 {NARRATIVE['mismatch_move_kj']} kJ。交叉验证闸门 {F['gate_status']}（月份差 {F['gate_delta_pct']:.1f}%），故 B、C 并报、主推仍报 B。化验加权约 {F['lab_mix']:.0f} kJ 对不上 {F['q_b']:.0f}——化验是样品，模型是入炉混合，不作为否决。", size=11, space=10)

    p(doc, "四、装修入炉怎么来的", size=14, bold=True, space=8, before=6)
    p(doc, "领导口径：造纸（其他）、农林按检测报告原值；装修用篮子剩余热反推。与综合数同一篮子（已剔克劳丽）。华衍沼渣单独钉自家化验，不并进装修。", size=11, space=8)
    table(doc, ["种类", "取值", "大卡/公斤", "吨", "占比"], [
        ["造纸（其他）理文底渣", "2026 年 2 月检测报告", "1626", "10,784", "16.1%"],
        ["农林 雷博尔", "2026 年 3 月检测报告", "2792", "9,582", "14.3%"],
        ["其他工业 东升", "2026 年 1 月检测报告", "892", "1,728", "2.6%"],
        ["华衍沼渣", "2026 年 5 月检测报告", "918", "3,223", "4.8%"],
        ["格栅 苏水中法", "2026 年 6 月检测报告", "1958", "532", "0.8%"],
        ["装修（反推）", "不钉化验，吃剩余的热", "约 800", "40,949", "61.3%"],
    ], highlight_last=True)
    p(doc, f"算法：篮子总热 = {F['q_b']:.0f} × {NARRATIVE['basket_t']:,} 吨；扣掉上表已钉种类后，除以装修吨。点估计 {NARRATIVE['decor_point_kj']} kJ/kg（{NARRATIVE['decor_point_kcal']} 大卡）。解法 C 对应 {NARRATIVE['decor_c_kcal']} 大卡。只钉造纸+农林、其余并进装修，得 {NARRATIVE['decor_pin2_kcal']} 大卡，与 {NARRATIVE['decor_point_kcal']} 几乎一样。", size=11, space=6)
    p(doc, f"天越、苏再投装修化验是 {NARRATIVE['lab_decor_kcal']} 大卡（筛上干样、偏干）。入炉是湿料、灰分、筛下不可燃、坑内掺混之后的反平衡。本算法承认：造纸、农林化验可以当入炉用，装修化验当不了。差在装修状态，不在综合数 {F['kcal_b']} 大卡。这不是分种类回归，误差会堆在装修上。", size=11, space=10)

    p(doc, "五、使用注意", size=14, bold=True, space=8, before=6)
    for line in [
        f"不要说「综合 {F['kcal_b']} 大卡，所以装修也是 {F['kcal_b']}」。",
        f"不要用装修化验 4000 多大卡否定入炉约 {NARRATIVE['decor_point_kcal']} 大卡。",
        "不要和 2025 年回归数直接比年（当年装修更重、无克劳丽，且内部验证未过，不能报价）。",
        "收到基低位、入炉态。不是干基，不是单车化验，不是商务报价单。",
    ]:
        p(doc, "·  " + line, size=11, space=3)
    p(doc, f"锁定：综合 {F['kcal_b']} 大卡；装修入炉约 {NARRATIVE['decor_point_kcal']} 大卡。", size=12, bold=True, space=8, before=10)
    p(doc, f"测算：固废智能运营台      数据时间：{F['generated']}", size=10, color=MUTED, space=0)

    fp = doc.sections[0].footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f"提交版  ·  综合 {F['kcal_b']} 大卡  ·  装修入炉约 {NARRATIVE['decor_point_kcal']} 大卡  ·  不含克劳丽")
    font(r, size=8, color=MUTED)
    doc.save(DOCX)
    print("word", DOCX)


def fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def box(slide, l, t, w, h, text, size=18, bold=False, color=C_INK, align=PP_ALIGN.LEFT):
    sh = slide.shapes.add_textbox(l, t, w, h)
    tf = sh.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size = PPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    return sh


def blank(prs, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    fill(bg, C_DARK if dark else C_BG)
    return s


def foot(slide, n, total=8):
    box(slide, Emu(500000), Emu(6420000), Emu(9000000), Emu(280000),
        "一般固废入炉热值  ·  提交版  ·  不含克劳丽", size=10, color=C_MUTE)
    box(slide, Emu(10500000), Emu(6420000), Emu(1200000), Emu(280000),
        f"{n} / {total}", size=10, color=C_MUTE, align=PP_ALIGN.RIGHT)


def card(slide, l, t, w, h):
    rec = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    rec.fill.solid()
    rec.fill.fore_color.rgb = C_WHITE
    rec.line.color.rgb = C_LINE
    return rec


def build_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    s = blank(prs, dark=True)
    box(s, Emu(600000), Emu(800000), Emu(11000000), Emu(320000),
        f"提交版  ·  {F['generated']}  ·  窗口 2 月 1 日～9 月 1 日", size=14, color=PptRGB(0xA8, 0xAD, 0xA8))
    box(s, Emu(600000), Emu(1250000), Emu(11000000), Emu(700000),
        "一般固废入炉热值", size=36, bold=True, color=C_WHITE)
    box(s, Emu(600000), Emu(2600000), Emu(5000000), Emu(350000), "综合固废  主报价", size=14, color=PptRGB(0xA8, 0xAD, 0xA8))
    box(s, Emu(600000), Emu(3000000), Emu(5000000), Emu(750000), f"{F['kcal_b']} 大卡", size=44, bold=True, color=C_WHITE)
    box(s, Emu(600000), Emu(3750000), Emu(5000000), Emu(320000), f"{F['q_b']:.0f} kJ/kg  ·  解法 B", size=15, color=PptRGB(0xA8, 0xAD, 0xA8))
    box(s, Emu(6400000), Emu(2600000), Emu(5000000), Emu(350000), "装修入炉  钉化验后残差", size=14, color=PptRGB(0xA8, 0xAD, 0xA8))
    box(s, Emu(6400000), Emu(3000000), Emu(5000000), Emu(750000), "约 800 大卡", size=44, bold=True, color=C_WHITE)
    box(s, Emu(6400000), Emu(3750000), Emu(5000000), Emu(320000), "3375 kJ/kg  ·  点估计 806", size=15, color=PptRGB(0xA8, 0xAD, 0xA8))
    box(s, Emu(600000), Emu(4700000), Emu(11000000), Emu(900000),
        "造纸（其他）、农林按检测报告取值，装修用篮子剩余热反推。\n已整日剔除克劳丽 18 天。内部验证通过。",
        size=16, color=PptRGB(0xC4, 0xC7, 0xC4))

    s = blank(prs)
    foot(s, 2)
    box(s, Emu(600000), Emu(380000), Emu(11000000), Emu(450000), "两个数怎么用", size=26, bold=True)
    for i, (t, b) in enumerate([
        (f"{F['kcal_b']} 大卡", f"谈这一篮子固废值多少热。相对生活垃圾约 {F['kcal_msw']} 大卡，大约是它的 {F['kcal_b']/F['kcal_msw']*100:.0f}%。区间约 {F['q_b_lo']/4.186:.0f}～{F['q_b_hi']/4.186:.0f} 大卡。"),
        (f"约 {NARRATIVE['decor_point_kcal']} 大卡", f"单独看装修。造纸底渣 1626 大卡、农林 2792 大卡按检测报告钉住。区间约 {NARRATIVE['decor_range_kcal']} 大卡，会上先报 {NARRATIVE['decor_point_kcal']}。"),
        ("不要混", f"{F['kcal_b']} 不是天越装修，也不是理文底渣。装修化验 {NARRATIVE['lab_decor_kcal']} 大卡是筛上干样，不是入炉。"),
    ]):
        top = 1050000 + i * 1650000
        card(s, Emu(600000), Emu(top), Emu(11000000), Emu(1500000)
        )
        box(s, Emu(850000), Emu(top + 180000), Emu(10500000), Emu(400000), t, size=20, bold=True, color=C_GREEN)
        box(s, Emu(850000), Emu(top + 650000), Emu(10500000), Emu(700000), b, size=16)

    s = blank(prs)
    foot(s, 3)
    box(s, Emu(600000), Emu(380000), Emu(11000000), Emu(450000), "数据与口径", size=26, bold=True)
    rows = [
        ("时间", "2026-02-01～09-01，213 日。历史尺子用 2024 二期纯烧，不用 2025。"),
        ("蒸汽", "生产日报一（流量）+ 指标日报（温压氧排烟）。只读，未改填报。"),
        ("固废", "白名单、货名含「其他」，跳过环卫所/城东。大件算生活垃圾。"),
        ("克劳丽", f"{NARRATIVE['kelali_days']} 日 / {NARRATIVE['kelali_cars']} 车 / {NARRATIVE['kelali_t']} 吨，整日剔除。计入则综合约 {NARRATIVE['with_kelali_kcal']} 大卡。"),
        ("样本", f"{F['n_weeks']} 周进回归。闸门 {F['gate_status']} {F['gate_delta_pct']:.1f}%，主报 B、并报 C。"),
    ]
    for i, (a, b) in enumerate(rows):
        top = 1000000 + i * 1000000
        box(s, Emu(600000), Emu(top), Emu(2000000), Emu(800000), a, size=16, bold=True, color=C_GREEN)
        box(s, Emu(2800000), Emu(top), Emu(8600000), Emu(850000), b, size=16)

    s = blank(prs)
    foot(s, 4)
    box(s, Emu(600000), Emu(380000), Emu(11000000), Emu(450000), "综合数怎么算", size=26, bold=True)
    box(s, Emu(600000), Emu(1000000), Emu(11000000), Emu(700000),
        "混料热值 = 生活 × (1 − α) + 固废 × α", size=22, bold=True)
    for i, t in enumerate([
        f"蒸汽热 ÷ 效率（{min(F['eta_p1'], F['eta_p2'])*100:.0f}%～{max(F['eta_p1'], F['eta_p2'])*100:.0f}%）÷ 入炉吨  →  这一锅有多热",
        "筛掉低负荷、缺测点、克劳丽日；一周不满 5 个有效日不进回归",
        f"α = 当周固废 / 二期入炉。{F['n_weeks']} 周从 {F['alpha_min_pct']:.0f}% 到 {F['alpha_max_pct']:.0f}%，均 {F['alpha_mean_pct']:.0f}%",
        f"一期当周当生活垃圾尺子，乘 k={F['k']:.2f} 换到二期表计，拧出固废 {F['q_b']:.0f}",
    ]):
        box(s, Emu(600000), Emu(1900000 + i * 1000000), Emu(11000000), Emu(850000),
            str(i + 1) + "    " + t, size=18)

    s = blank(prs)
    foot(s, 5)
    box(s, Emu(600000), Emu(350000), Emu(11000000), Emu(420000), "报 B，并报 C；验证通过", size=24, bold=True)
    trips = [
        ("解法 B  主推", f"{F['kcal_b']} 大卡", "一期同周 × k"),
        ("解法 C  并报", f"{F['kcal_c']} 大卡", f"与 B 差 {abs(F['kcal_b']-F['kcal_c'])} 大卡"),
        ("解法 A  校验", f"{F['kcal_a']} 大卡", "区间宽，不单报"),
    ]
    for i, (h, n, f) in enumerate(trips):
        left = 500000 + i * 3800000
        card(s, Emu(left), Emu(950000), Emu(3550000), Emu(2500000))
        box(s, Emu(left + 180000), Emu(1100000), Emu(3200000), Emu(400000), h, size=14, color=C_GREEN)
        box(s, Emu(left + 180000), Emu(1550000), Emu(3200000), Emu(900000), n, size=26, bold=True)
        box(s, Emu(left + 180000), Emu(2600000), Emu(3200000), Emu(500000), f, size=14, color=C_MUTE)
    for i, t in enumerate([
        f"V0  B 与 C 差 {F['b_minus_c']:.0f} kJ，门限 500。",
        "V2  高掺烧周仍落在区间内。",
        f"V3  踢掉残差周后只挪 {NARRATIVE['mismatch_move_kj']} kJ。",
        f"化验 {F['lab_mix']:.0f} 对不上 {F['q_b']:.0f}，不否决（样品 ≠ 入炉）。",
    ]):
        box(s, Emu(600000), Emu(3700000 + i * 600000), Emu(11000000), Emu(550000), t, size=16)

    s = blank(prs)
    foot(s, 6)
    box(s, Emu(600000), Emu(300000), Emu(11000000), Emu(420000), "装修：钉检测报告，吃剩余的热", size=24, bold=True)
    headers = ["种类", "大卡", "吨", "占比"]
    data = [
        ["造纸（其他）理文底渣", "1626", "10,784", "16%"],
        ["农林 雷博尔", "2792", "9,582", "14%"],
        ["其他工业 / 沼渣 / 格栅", "892～1958", "5,483", "8%"],
        ["装修（反推，不钉化验）", "约 800", "40,949", "61%"],
    ]
    col_w = [4800000, 2000000, 2000000, 1800000]
    lefts = [600000]
    for w in col_w[:-1]:
        lefts.append(lefts[-1] + w)
    top0, rh = 850000, 900000
    for ci, h in enumerate(headers):
        cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(lefts[ci]), Emu(top0), Emu(col_w[ci]), Emu(rh))
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_DARK
        cell.line.fill.background()
        box(s, Emu(lefts[ci] + 100000), Emu(top0 + 250000), Emu(col_w[ci] - 160000), Emu(400000),
            h, size=14, bold=True, color=C_WHITE)
    for ri, row in enumerate(data):
        top = top0 + (ri + 1) * rh
        bg = C_WHITE if ri < 3 else PptRGB(0xED, 0xF3, 0xEC)
        for ci, val in enumerate(row):
            cell = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(lefts[ci]), Emu(top), Emu(col_w[ci]), Emu(rh))
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cell.line.color.rgb = C_LINE
            box(s, Emu(lefts[ci] + 100000), Emu(top + 250000), Emu(col_w[ci] - 160000), Emu(400000),
                val, size=15, bold=(ri == 3))
    box(s, Emu(600000), Emu(5480000), Emu(11000000), Emu(700000),
        "点估计 806 大卡。只钉造纸+农林得 830 大卡，几乎一样。", size=15, color=C_MUTE)

    s = blank(prs)
    foot(s, 7)
    box(s, Emu(600000), Emu(380000), Emu(11000000), Emu(450000), "化验装修不是入炉装修", size=24, bold=True)
    card(s, Emu(600000), Emu(1100000), Emu(5200000), Emu(3200000))
    box(s, Emu(850000), Emu(1300000), Emu(4700000), Emu(400000), "化验室  筛上干样", size=14, color=C_MUTE)
    box(s, Emu(850000), Emu(1850000), Emu(4700000), Emu(800000), "4438～5747 大卡", size=26, bold=True)
    box(s, Emu(850000), Emu(2800000), Emu(4700000), Emu(1100000), "天越 7 月、苏再投 4 月\n拣过的轻质，偏干", size=15, color=C_MUTE)
    card(s, Emu(6100000), Emu(1100000), Emu(5200000), Emu(3200000))
    box(s, Emu(6350000), Emu(1300000), Emu(4700000), Emu(400000), "入炉  湿料残差", size=14, color=C_MUTE)
    box(s, Emu(6350000), Emu(1850000), Emu(4700000), Emu(800000), "约 800 大卡", size=26, bold=True)
    box(s, Emu(6350000), Emu(2800000), Emu(4700000), Emu(1100000), "含水、灰、筛下物、坑内掺混\n造纸/农林化验可当入炉用", size=15, color=C_MUTE)
    box(s, Emu(600000), Emu(4550000), Emu(11000000), Emu(1400000),
        "差在装修状态，不在综合数。误差堆在装修上，不是分种类回归。\n克劳丽已剔除：计入则综合约 1309 大卡，装修会被它往上拉。",
        size=16)

    s = blank(prs, dark=True)
    box(s, Emu(600000), Emu(1200000), Emu(11000000), Emu(400000), "锁定", size=16, color=PptRGB(0xA8, 0xAD, 0xA8))
    box(s, Emu(600000), Emu(1700000), Emu(11000000), Emu(900000),
        f"综合 {F['kcal_b']} 大卡      装修入炉约 {NARRATIVE['decor_point_kcal']} 大卡", size=28, bold=True, color=C_WHITE)
    box(s, Emu(600000), Emu(2900000), Emu(11000000), Emu(1800000),
        "不要把两个数混成一个。不要用装修干样化验否定入炉数。\n不要和 2025 年直接比年。不是商务报价单。",
        size=18, color=PptRGB(0xC4, 0xC7, 0xC4))

    try:
        prs.save(PPTX)
        print("ppt", PPTX)
    except PermissionError:
        alt = OUT / "固废热值测算说明_提交版_中和.pptx"
        prs.save(alt)
        print("ppt locked, wrote", alt)


if __name__ == "__main__":
    build_docx()
    build_pptx()
