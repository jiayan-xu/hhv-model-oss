# -*- coding: utf-8 -*-
"""2025 日表：会话导出的生产日报一（产汽）+ 指标日报温压 + 地磅分类固废。不覆盖 2026 CSV。"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "snmis_raw"
OUT = ROOT / "data" / "plant" / "y2025"
CARGO = ROOT / "data" / "plant" / "cargo_2025_typed.csv"

COL = dict(mass=4, steam_t=7, steam_p=10, feedwater_t=13, evap=16, fluegas_t=22, o2=25)
BOILER_ROWS = (20, 21, 22)


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def parse_file(path: Path) -> dict | None:
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    recs = []
    evap_sum = 0.0
    for r in BOILER_ROWS:
        mass = _num(ws.cell(r, COL["mass"]).value) or 0.0
        p = _num(ws.cell(r, COL["steam_p"]).value)
        t = _num(ws.cell(r, COL["steam_t"]).value)
        evap = _num(ws.cell(r, COL["evap"]).value) or 0.0
        if p is None or t is None or p < 1.0 or t < 200:
            continue
        w = mass if mass > 0 else evap
        if w <= 0:
            w = 1.0
        recs.append({
            "w": w,
            "steam_p": p,
            "steam_t": t,
            "feedwater_t": _num(ws.cell(r, COL["feedwater_t"]).value),
            "o2": _num(ws.cell(r, COL["o2"]).value),
            "fluegas_t": _num(ws.cell(r, COL["fluegas_t"]).value),
        })
        evap_sum += evap
    if not recs:
        return None
    tw = sum(x["w"] for x in recs)
    out = {"n_on": len(recs), "phase_input": tw, "steam_t_d": evap_sum,
           "steam_flow": evap_sum / 24.0}
    for k in ("steam_p", "steam_t", "feedwater_t", "o2", "fluegas_t"):
        nums = [(x["w"], x[k]) for x in recs if x[k] is not None]
        out[k] = (sum(w * v for w, v in nums) / sum(w for w, _ in nums)) if nums else None
    out["plant_in"] = _num(ws.cell(5, 7).value) or 0.0
    out["plant_q"] = _num(ws.cell(7, 7).value)
    out["kitchen"] = 0.0
    return out


def collect(glob: str) -> pd.DataFrame:
    rows = []
    for p in sorted(RAW.glob(glob)):
        iso = p.stem.split("_", 2)[-1]
        rec = parse_file(p)
        if rec is None:
            continue
        rec["time"] = iso
        rows.append(rec)
    df = pd.DataFrame(rows)
    if df.empty:
        return df
    df["time"] = pd.to_datetime(df["time"])
    return df.sort_values("time").drop_duplicates("time")


def overlay_tp(df: pd.DataFrame, old_path: Path) -> pd.DataFrame:
    """日报产汽 + 指标日报温压（旧 y2025 文件已从 meeting 抽过）。"""
    if not old_path.exists():
        return df
    old = pd.read_csv(old_path, parse_dates=["time"])
    keys = ["steam_p", "steam_t", "feedwater_t", "o2", "fluegas_t"]
    m = df.merge(old[["time"] + keys], on="time", how="left", suffixes=("", "_m"))
    for k in keys:
        mk = k + "_m"
        if mk not in m.columns:
            continue
        hit = m[mk].notna()
        m.loc[hit, k] = m.loc[hit, mk]
        m.drop(columns=[mk], inplace=True)
    return m


def main() -> int:
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        "bsd", ROOT / "scripts" / "build_snmis_daily.py")
    bsd = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(bsd)

    OUT.mkdir(parents=True, exist_ok=True)
    old1, old2 = OUT / "phase1_daily.csv", OUT / "phase2_daily.csv"
    p1, p2 = bsd.collect_year(2025)
    p1 = overlay_tp(p1, old1)
    p2 = overlay_tp(p2, old2)
    cols = ["time", "steam_flow", "steam_p", "steam_t", "feedwater_t", "o2", "fluegas_t"]
    p1[cols].to_csv(OUT / "phase1_daily.csv", index=False, encoding="utf-8-sig")
    p2[cols].to_csv(OUT / "phase2_daily.csv", index=False, encoding="utf-8-sig")

    cargo = pd.read_csv(CARGO, parse_dates=["date"])
    cargo["date"] = pd.to_datetime(cargo["date"]).dt.normalize()
    cargo["isw"] = (
        cargo["isw_装修"].fillna(0) + cargo["isw_农林"].fillna(0)
        + cargo["isw_底渣"].fillna(0) + cargo["isw_格栅"].fillna(0)
        + cargo["isw_其他工业"].fillna(0)
    )
    led = pd.DataFrame({"date": p1["time"].dt.normalize()})
    led["phase1_input"] = p1["phase_input"].values
    led["kitchen"] = p1["kitchen"].values
    led["p1_in"] = p1["plant_in"].values
    led["plant_q_p1"] = p1["plant_q"].values
    p2c = p2[["time", "phase_input", "plant_in", "plant_q"]].rename(columns={
        "time": "date", "phase_input": "phase2_input", "plant_in": "p2_in",
        "plant_q": "plant_q_p2",
    })
    p2c["date"] = pd.to_datetime(p2c["date"]).dt.normalize()
    led = led.merge(p2c, on="date", how="outer")
    cg = cargo[["date", "isw", "msw"]].rename(columns={"msw": "msw_in"})
    led = led.merge(cg, on="date", how="left")
    led["isw"] = led["isw"].fillna(0.0)
    fallback = (led["p1_in"].fillna(0) + led["p2_in"].fillna(0) - led["isw"]).clip(lower=0)
    led["msw_in"] = led["msw_in"].where(led["msw_in"].fillna(0) > 0, fallback)
    led["kitchen"] = led["kitchen"].fillna(0.0)
    keep = ["date", "phase1_input", "phase2_input", "msw_in", "isw",
            "p1_in", "p2_in", "kitchen", "plant_q_p1", "plant_q_p2"]
    led[keep].to_csv(OUT / "ledger_daily.csv", index=False, encoding="utf-8-sig")
    print("p1", len(p1), p1["time"].min().date(), p1["time"].max().date(),
          "steam med", round(float(p1["steam_flow"].median()), 2))
    print("p2", len(p2), "steam med", round(float(p2["steam_flow"].median()), 2))
    print("吨汽 p1", round(24 * float(p1["steam_flow"].median()) / float(led["phase1_input"].median()), 2),
          "p2", round(24 * float(p2["steam_flow"].median()) / float(led["phase2_input"].median()), 2))
    print("ledger", len(led), "isw", round(float(led["isw"].sum()), 1),
          "msw", round(float(led["msw_in"].sum()), 1))
    print("wrote", OUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
