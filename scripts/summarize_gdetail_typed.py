# -*- coding: utf-8 -*-
"""按 dashboard VE 口径拆地磅日明细种类（货名含「其他」+ 跳过环卫所/城东 + 白名单）。

分类来源：B 列括号（装修/农林）+ mappings.UNIT_WASTE_TYPE。
不覆盖 cargo_daily.csv。输出 data/plant/cargo_YYYY_typed.csv。
"""
from __future__ import annotations

import re
import sqlite3
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "data" / "snmis_raw" / "gdetail"
OUT = ROOT / "data" / "plant"
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
from _local_env import dashboard_db  # noqa: E402
DASH = dashboard_db().parent
sys.path.insert(0, str(DASH))

from mappings import (  # noqa: E402
    SKIP_PLATES,
    SKIP_UNITS,
    UNIT_WASTE_TYPE,
    normalize_company,
    normalize_waste,
)

MSW_GOODS = {
    "生活源垃圾": "msw_生活",
    "昆山生活源垃圾": "msw_昆山",
    "厨余垃圾": "msw_厨余",
    "餐厨垃圾": "msw_餐厨",
    "机扫垃圾": "msw_机扫",
    "其他（大件粉碎垃圾）": "msw_大件",
    "陈腐垃圾": "msw_陈腐",
}
TRACK_GOODS = {
    "秸秆": "track_秸秆",
    "固渣": "track_固渣",
    "装修轻质物": "track_货名装修轻质物",
}


def _num(v) -> float:
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _bucket(wt: str) -> str:
    if not wt:
        return "未映射"
    if "农林" in wt:
        return "农林"
    if "装修" in wt:
        return "装修"
    if "底渣" in wt:
        return "底渣"
    if "格栅" in wt:
        return "格栅"
    if "其他" in wt or "工业" in wt:
        return "其他工业"
    return wt


_TAG_RE = re.compile(r"[\(（]([^）)]{1,12})[\)）]")


def _split_unit(area: str) -> tuple[str, str | None]:
    """B 列单位 + 半角/全角括号标注。"""
    m = _TAG_RE.search(area)
    if not m:
        return area.strip(), None
    return area[: m.start()].strip(), m.group(1).strip()


def _waste_type(unit: str, tag: str | None) -> str:
    if tag:
        if "农林" in tag:
            return "农林垃圾"
        if "装修" in tag:
            return "装修垃圾轻质筛分物"
        if "底渣" in tag:
            return "底渣SW15"
        if "格栅" in tag:
            return "格栅垃圾SW59"
    if not unit:
        return ""
    if "苏州天越" in unit or "利合" in unit:
        return "装修垃圾轻质筛分物"
    for key, waste_type in UNIT_WASTE_TYPE.items():
        if key in unit:
            return waste_type
    return ""


def _whitelist() -> set[str]:
    db = DASH / "dashboard.db"
    if not db.exists():
        return set()
    conn = sqlite3.connect(db)
    rows = conn.execute(
        "SELECT license_plate FROM vehicle_whitelist WHERE enabled=1"
    ).fetchall()
    conn.close()
    return {r[0] for r in rows}


def parse_file(path: Path, plates: set[str]) -> dict[str, float]:
    wb = load_workbook(path, data_only=True)
    acc: dict[str, float] = defaultdict(float)
    last_u = None
    for sn in wb.sheetnames:
        ws = wb[sn]
        phase = "p1" if "一" in sn else "p2"
        last_u = None
        for r in range(3, ws.max_row + 1):
            raw_u = ws.cell(r, 2).value
            if raw_u is not None and str(raw_u).strip():
                last_u = raw_u
            g = ws.cell(r, 5).value
            if g is None:
                continue
            gs = str(g).strip()
            if not gs or gs == "/" or "合计" in gs:
                continue
            w = _num(ws.cell(r, 9).value)
            if w <= 0:
                continue
            acc[f"{phase}_net"] += w
            if gs in MSW_GOODS:
                acc[MSW_GOODS[gs]] += w
                acc["msw"] += w
                continue
            if gs in TRACK_GOODS:
                acc[TRACK_GOODS[gs]] += w
                continue
            if "其他" not in gs:
                acc["other_unmapped"] += w
                acc[f"raw::{gs}"] += w
                continue

            acc["isw_goods_其他"] += w
            unit_s = str(last_u or "").strip()
            if "合计" in unit_s:
                continue
            if "城东" in unit_s:
                acc["drop_城东"] += w
                continue
            unit_name, tag = _split_unit(unit_s)
            if any(kw in unit_name for kw in SKIP_UNITS):
                acc["drop_环卫所"] += w
                continue
            plate = str(ws.cell(r, 4).value or "").strip()
            if plate in SKIP_PLATES:
                continue
            unit_name = normalize_company(unit_name)
            wt = normalize_waste(_waste_type(unit_name, tag))
            bucket = _bucket(wt)
            acc["isw_skip环卫"] += w
            acc[f"isw_{bucket}"] += w
            if plates and plate not in plates:
                acc["drop_非白名单"] += w
                acc[f"drop_unit::{unit_name}"] += w
                continue
            acc["isw_ve"] += w
            acc[f"ve_{bucket}"] += w
            acc[f"ve_unit::{unit_name}"] += w
    acc["date"] = path.stem.replace("gdetail_", "")
    return acc


def main() -> int:
    year = int(sys.argv[1]) if len(sys.argv) > 1 else 2025
    files = sorted(SRC.glob(f"gdetail_{year}-*.xlsx"))
    if not files:
        raise FileNotFoundError(f"{SRC}/gdetail_{year}-*.xlsx")
    plates = _whitelist()
    rows = [parse_file(p, plates) for p in files]
    df = pd.DataFrame(rows)
    df["date"] = pd.to_datetime(df["date"])
    keep = [
        "date",
        "isw_ve",
        "ve_装修",
        "ve_农林",
        "ve_底渣",
        "ve_格栅",
        "ve_其他工业",
        "ve_未映射",
        "isw_装修",
        "isw_农林",
        "isw_底渣",
        "isw_格栅",
        "isw_其他工业",
        "isw_未映射",
        "isw_goods_其他",
        "isw_skip环卫",
        "drop_环卫所",
        "drop_城东",
        "drop_非白名单",
        "msw",
        "msw_生活",
        "msw_昆山",
        "msw_厨余",
        "msw_餐厨",
        "msw_机扫",
        "msw_大件",
        "msw_陈腐",
        "track_秸秆",
        "track_固渣",
        "track_货名装修轻质物",
        "other_unmapped",
        "p1_net",
        "p2_net",
    ]
    for c in keep:
        if c not in df.columns:
            df[c] = 0.0
    daily = df[keep].fillna(0.0).sort_values("date")
    OUT.mkdir(parents=True, exist_ok=True)
    dest = OUT / f"cargo_{year}_typed.csv"
    daily.to_csv(dest, index=False, encoding="utf-8-sig")

    m = daily.copy()
    m["month"] = m["date"].dt.to_period("M").astype(str)
    num = [c for c in keep if c != "date"]
    monthly = m.groupby("month")[num].sum().round(1)
    md = OUT / f"cargo_{year}_typed_monthly.csv"
    monthly.to_csv(md, encoding="utf-8-sig")

    print(dest, "days", len(daily), daily["date"].min().date(), daily["date"].max().date())
    print("whitelist", len(plates))
    print("\n=== 年合计 t ===")
    tot = daily[num].sum().round(1)
    show = [
        "isw_ve", "ve_装修", "ve_农林", "ve_底渣", "ve_格栅", "ve_其他工业", "ve_未映射",
        "isw_skip环卫", "isw_装修", "isw_农林", "isw_底渣", "isw_格栅", "isw_其他工业", "isw_未映射",
        "isw_goods_其他", "drop_环卫所", "drop_城东", "drop_非白名单",
        "msw", "msw_生活", "msw_昆山", "msw_厨余", "msw_大件", "msw_陈腐",
        "track_固渣", "track_货名装修轻质物",
    ]
    print(tot[show].to_string())
    print("\n=== 月 全量「其他」(跳过环卫/城东) 分种类 t ===")
    print(monthly[["isw_skip环卫", "isw_装修", "isw_农林", "isw_底渣", "isw_格栅", "isw_其他工业", "isw_未映射"]].to_string())
    print("\n=== 月 VE(现白名单) 分种类 t ===")
    print(monthly[["isw_ve", "ve_装修", "ve_农林", "ve_底渣", "ve_格栅", "ve_其他工业", "ve_未映射"]].to_string())

    units: Counter[str] = Counter()
    dropped: Counter[str] = Counter()
    unmapped: Counter[str] = Counter()
    for _, row in df.iterrows():
        for c, v in row.items():
            if not isinstance(c, str) or v is None:
                continue
            try:
                fv = float(v)
            except (TypeError, ValueError):
                continue
            if fv != fv or fv <= 0:
                continue
            if c.startswith("ve_unit::"):
                units[c.replace("ve_unit::", "")] += fv
            elif c.startswith("drop_unit::"):
                dropped[c.replace("drop_unit::", "")] += fv
            elif c.startswith("raw::"):
                unmapped[c.replace("raw::", "")] += fv
    print("\n=== VE 企业 top ===")
    for k, v in units.most_common(20):
        print(f"  {v:8.1f}  {k}")
    print("\n=== 非白名单企业 top（货名含其他） ===")
    for k, v in dropped.most_common(12):
        print(f"  {v:8.1f}  {k}")
    if unmapped:
        print("\n=== 货名未映射 ===")
        for k, v in unmapped.most_common(15):
            print(f"  {v:8.1f}  {k}")
    print("\nmonthly", md)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
