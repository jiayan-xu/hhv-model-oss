# -*- coding: utf-8 -*-
"""按模版重生成 固废热值测算_20260903.pptx（干净包，不补丁旧文件）。

内容页 1～4 与原模版一致；原「定价接入概况」换成报告第六部分：
6.1 厂内口径 / 6.2 0费单批次 / 6.3 装修分档 / 6.4 农林一半。
不含原 6.3 吨位账、不含 6.6 掺烧 50%。
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

OUT = Path(__file__).resolve().parents[1] / "outputs" / "config.snmis_daily" / "固废热值测算_20260903.pptx"

FONT = "微软雅黑"
INK = RGBColor(0x17, 0x21, 0x3A)
MUTE = RGBColor(0x6B, 0x7C, 0x93)
BLUE = RGBColor(0x1E, 0x88, 0xE5)
GREEN = RGBColor(0x17, 0x8A, 0x50)
ORANGE = RGBColor(0xD9, 0x77, 0x06)
NAVY = RGBColor(0x17, 0x21, 0x3A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PILL = RGBColor(0xEE, 0xF2, 0xF8)
BG = RGBColor(0xF7, 0xF9, 0xFC)
TOTAL = 9
W, H = Emu(12191365), Emu(6858000)


def add_blank(prs, navy=False):
  s = prs.slides.add_slide(prs.slide_layouts[6])
  s.background.fill.solid()
  s.background.fill.fore_color.rgb = NAVY if navy else BG
  return s


def run_style(run, size, bold, color):
  run.font.size = Pt(size)
  run.font.bold = bold
  run.font.color.rgb = color
  run.font.name = FONT


def tb(slide, l, t, w, h, lines, *, fill=None, align=PP_ALIGN.LEFT):
  sh = slide.shapes.add_textbox(Emu(l), Emu(t), Emu(w), Emu(h))
  tf = sh.text_frame
  tf.word_wrap = True
  if fill is not None:
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    sh.line.fill.background()
  for i, item in enumerate(lines):
    text, size, bold, color = item
    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    para.alignment = align
    para.space_after = Pt(2)
    run = para.add_run()
    run.text = text
    run_style(run, size, bold, color)
  return sh


def header(slide, n, kicker="一般固废入炉热值 · 池内混料模型"):
  tb(slide, 548640, 320040, 9000000, 260350, [(kicker, 11, False, MUTE)])
  tb(slide, 10789920, 320040, 914400, 260350,
     [(f"{n} / {TOTAL}", 11, False, MUTE)], align=PP_ALIGN.RIGHT)


def cover(prs):
  s = add_blank(prs)
  tb(s, 548640, 457200, 10972800, 275590,
     [("2026-09-03  ·  窗口 2 月 1 日～9 月 1 日", 12, False, MUTE)])
  tb(s, 822960, 1371600, 10515600, 1097280,
     [("一般固废入炉热值", 40, True, INK)])
  tb(s, 822960, 2743200, 5029200, 2011680, [
    ("综合固废  主报价", 15, False, MUTE),
    ("1188 大卡", 44, True, BLUE),
    ("4972 kJ/kg · 解法 B", 14, False, INK),
    ("两个独立锚相互验证：差仅 43 kJ（10 大卡）", 12, False, GREEN),
  ])
  tb(s, 6400800, 2743200, 5029200, 2011680, [
    ("装修入炉  钉化验后残差", 15, False, MUTE),
    ("约 720 大卡", 44, True, ORANGE),
    ("3016 kJ/kg · 点估计", 14, False, INK),
    ("造纸、农林按检测报告钉住，装修吃剩余", 12, False, MUTE),
  ])
  tb(s, 822960, 5303520, 10515600, 700000, [
    ("垃圾池混料物理模型 + 化验值固定 + 装修残差反推。已整日剔除克劳丽。", 13, False, INK),
    ("修正了回归变量噪声偏置，两锚收敛更强。旧数在新区间内。", 12, False, MUTE),
  ])


def slide_koujing(prs):
  s = add_blank(prs)
  header(s, 1)
  tb(s, 822960, 914400, 10515600, 548640, [("数据与口径", 26, True, INK)])
  rows = [
    ("时间", "2026-02-01～09-01，213 日。历史尺子用 2024 年二期纯烧生活垃圾历史数据。"),
    ("蒸汽", "SNMIS 系统导出：生产日报一（流量）+ 指标日报（温压氧排烟）。"),
    ("固废", "固废监管系统中所定义的一般固废。"),
    ("克劳丽", "18 日 / 387 车 / 8124 吨，属于偶发扰动事件，整日剔除。"),
    ("垃圾池", "滞后按数据选（扫描 0～7 天取最优），α 取进厂成分 3 周滚动均值——混料池物理真实。"),
    ("样本", "24 周进回归。闸门 DUAL 9.5%，主报 B、并报 C。"),
  ]
  y0, dy = 1828800, 712000
  for i, (lab, val) in enumerate(rows):
    y = y0 + i * dy
    tb(s, 822960, y, 1828800, 548640, [(lab, 14, True, BLUE)])
    tb(s, 2743200, y, 8686800, 620000, [(val, 13, False, INK)])


def slide_gongshi(prs):
  s = add_blank(prs)
  header(s, 2)
  tb(s, 822960, 685800, 10515600, 400000, [("综合数怎么算：公式与推导", 24, True, INK)])
  tb(s, 822960, 1120000, 10515600, 280000, [("第一步  锅炉反平衡 → 混料热值 Q混", 15, True, BLUE)])
  tb(s, 822960, 1420000, 5852160, 1000000, [
    ("Q入炉 = D×(h汽−h水) / η", 13, False, INK),
    ("η = 1 − (q₂+q₃+q₄+q₅+q₆)", 12, False, INK),
    ("Q混 = Q入炉 / M入炉   (kJ/kg)", 12, False, INK),
  ], fill=PILL)
  tb(s, 6949440, 1420000, 4572000, 1100000, [
    ("D=蒸汽流量，h=焓差，η=反平衡效率（86~87%），M=入炉吨。这一步得到「这一锅有多热」。", 12, False, MUTE),
  ])
  tb(s, 822960, 2680000, 10515600, 280000, [("第二步  一期纯烧当尺子 → Q生活（锚）", 15, True, BLUE)])
  tb(s, 822960, 2980000, 5852160, 1000000, [
    ("Q生活锚(w) = Q生活一期(w) × k", 13, False, INK),
    ("k = median(Q二期历史 / Q一期当月)", 12, False, INK),
    ("本窗口 k = 1.10（表计刻度修正）", 12, False, INK),
  ], fill=PILL)
  tb(s, 6949440, 2980000, 4572000, 1100000, [
    ("一期几乎纯烧生活垃圾；二期两年前有纯烧历史段，同月份交叉验证求出两炉表计刻度比 k。", 12, False, MUTE),
  ])
  tb(s, 822960, 4240000, 10515600, 280000, [("第三步  掺烧比例分解 → 回归解出 Q固废", 15, True, BLUE)])
  tb(s, 822960, 4540000, 5852160, 1100000, [
    ("Q混(w) = Q生活(w)×(1−α) + Q固废×α", 13, False, INK),
    ("α(w) = m固废(w) / M入炉(w)  ← 3周平滑", 12, False, INK),
    ("Huber 稳健回归 → Q固废 = 4972 kJ/kg", 12, False, INK),
  ], fill=PILL)
  tb(s, 6949440, 4540000, 4572000, 1100000, [
    ("α 取进厂成分 3 周滚动均值；Bootstrap 2000 次给 95% 区间 4324~5700。", 12, False, MUTE),
  ])
  tb(s, 822960, 5800000, 10515600, 500000, [
    ("核心改进：α 池内平滑后，独立双锚（B/C）收敛到差 43 kJ——公式没变，回归变量噪声消了。", 12, False, GREEN),
  ])


def slide_zhuangxiu(prs):
  s = add_blank(prs)
  header(s, 3)
  tb(s, 822960, 800000, 10515600, 420000, [("装修：钉检测报告，吃剩余的热", 26, True, INK)])
  heads = ["种类", "大卡", "吨", "占比"]
  xs = [822960, 3474720, 6126480, 8778240]
  for x, h in zip(xs, heads):
    tb(s, x, 1320000, 2468880, 360000, [(h, 13, True, BLUE)])
  rows = [
    ("造纸（其他）理文底渣", "1626", "10,784", "16%", INK, False),
    ("农林 雷博尔", "2792", "9,582", "14%", INK, False),
    ("其他工业 / 沼渣 / 格栅", "892～1958", "5,483", "8%", INK, False),
    ("装修（反推，不钉化验）", "约 720", "40,949", "61%", ORANGE, True),
  ]
  for i, (a, b, c, d, col, bold) in enumerate(rows):
    y = 1750000 + i * 520000
    for x, val in zip(xs, (a, b, c, d)):
      tb(s, x, y, 2468880, 480000, [(val, 14, bold, col)])
  tb(s, 822960, 4000000, 10515600, 1100000, [
    ("点估计 720 大卡。只钉造纸+农林得约 856 大卡，敏感性稳定。", 13, False, INK),
    ("装修占吨位 61%，综合数的区间传到装修被放大——这是方法边界，不是数据错误。", 12, False, MUTE),
  ])


def slide_zhuashou(prs):
  s = add_blank(prs)
  header(s, 4)
  tb(s, 822960, 800000, 10515600, 420000, [("模型精度还能提多少：三个抓手", 26, True, INK)])
  heads = ["措施", "投入", "预期收益", "状态"]
  xs = [822960, 4480560, 6675120, 10515600]
  ws = [3520440, 2057400, 3703320, 1600000]
  for x, w, h in zip(xs, ws, heads):
    tb(s, x, 1320000, w, 400000, [(h, 12, True, BLUE)])
  rows = [
    ("① 装修入炉湿样化验 2~3 批", "约几百元/样", "装修从纯残差→有独立锚；装修 CI 至少砍半", "建议尽快做", GREEN),
    ("② 垃圾池库存数据接入", "目测料位 30秒/天", "α 噪声再降；综合 CI 再收 1~2 个点", "方案已备，待启动", ORANGE),
    ("③ DCS 小时数据拉取（合格周）", "约 700 次请求，挂机一天", "炉侧计量精度提升；效率 η 更准", "接口已验证，待批量", ORANGE),
  ]
  for i, (a, b, c, d, col) in enumerate(rows):
    y = 1800000 + i * 850000
    vals = (a, b, c, d)
    colors = (INK, INK, INK, col)
    for x, w, val, color in zip(xs, ws, vals, colors):
      tb(s, x, y, w, 800000, [(val, 12, False, color)])
  tb(s, 822960, 4450000, 10515600, 1200000, [
    ("优先级：①最值（装修单独定价的直接依据）→ ②零成本（方案已写好）→ ③锦上添花。", 13, False, INK),
    ("三项全做后综合 CI 有望从 ±14% 收到 ±8~10%，装修从「区间估计」升级为「有锚测算」。", 13, False, GREEN),
  ])


def slide_61(prs):
  s = add_blank(prs)
  header(s, 5, "一般固废入炉热值 · 定价模型接入")
  tb(s, 822960, 780000, 10515600, 400000, [("厂内口径", 26, True, INK)])
  rows = [
    ("入炉态", "反推已是入炉热值，堆放天数=0，不再打堆放折"),
    ("含水", "厂内 k_moist=1，检测低位已是该含水收到基"),
    ("炉渣运出价", "30 元/吨渣（出厂卖价）"),
    ("炉渣率", "出厂 127,684 t / 入炉 514,380 t = 24.82%"),
    ("摊到垃圾", "30 × 24.82% = 7.5 元/t，不是每吨垃圾进账 30"),
    ("掺烧档", "按 20% 档取效率、厂用电与耗材"),
    ("变动成本已含", "石灰、活性炭、氨水、天然气、飞灰、检修、渗滤液、除盐水"),
    ("未进变动成本", "折旧、运行人工、财务费用（吨固定约 47 元）"),
  ]
  y0, dy = 1220000, 580000
  for i, (lab, val) in enumerate(rows):
    y = y0 + i * dy
    tb(s, 822960, y, 2200000, 520000, [(lab, 14, True, BLUE)])
    tb(s, 3100000, y, 8200000, 520000, [(val, 13, False, INK)])


def slide_62(prs):
  s = add_blank(prs)
  header(s, 6, "一般固废入炉热值 · 定价模型接入")
  tb(s, 822960, 740000, 10515600, 380000,
     [("单批次 · 处置费=0 · 变动成本层", 24, True, INK)])
  heads = ["批次", "入炉大卡", "发电", "炉渣", "变动成本", "0 费边际", "盈亏平衡 B"]
  data = [
    ("综合固废", "1188", "96.3", "7.5", "80.1", "+23.6", "−23.6"),
    ("装修", "720", "58.4", "7.5", "77.1", "−11.3", "+11.3"),
    ("造纸（理文）", "1626", "131.8", "7.5", "88.6", "+50.6", "−50.6"),
    ("农林（雷博尔）", "2792", "226.3", "7.5", "78.0", "+155.8", "−155.8"),
  ]
  xs = [822960, 2500000, 4000000, 5300000, 6600000, 8300000, 10000000]
  ws = [1600000, 1400000, 1200000, 1200000, 1600000, 1600000, 1400000]
  for x, w, h in zip(xs, ws, heads):
    tb(s, x, 1180000, w, 340000, [(h, 12, True, BLUE)])
  for ri, row in enumerate(data):
    y = 1560000 + ri * 480000
    color = ORANGE if row[0] == "装修" else INK
    bold = row[0] == "装修"
    for x, w, cell in zip(xs, ws, row):
      tb(s, x, y, w, 440000, [(cell, 14, bold, color)])
  tb(s, 822960, 3600000, 10515600, 1800000, [
    ("单位：元/t。B = 变动成本 − 发电 − 炉渣。B>0 表示 0 处置费时变动成本层亏损。", 13, False, INK),
    ("综合 1188 大卡发电盖住药剂，0 费仍 +23.6。装修 720 不够热，药费按吨走，0 费 −11.3。打平约 860～900 大卡。", 13, False, INK),
    ("全成本另说：吨固定约 47 元。综合 0 费扣折旧人工后约 −24。药剂电渣这层不亏，摊上固定成本仍亏。", 12, False, MUTE),
    ("变动成本已含石灰、活性炭、氨水、天然气（约 1.17 元/t）、飞灰 40、检修 16。草图，不是合同价。", 12, False, MUTE),
  ])


def slide_63(prs):
  s = add_blank(prs)
  header(s, 7, "一般固废入炉热值 · 定价模型接入")
  tb(s, 822960, 740000, 10515600, 380000,
     [("装修处置价分档  ·  σ=15%", 24, True, INK)])
  heads = ["处置价 元/t", "期望边际", "P5 边际", "亏损概率", "判定"]
  data = [
    ("0", "−11.3", "−25.7", "90%", "高风险", RGBColor(0xC0, 0x39, 0x2B)),
    ("10", "−1.3", "−15.7", "56%", "高风险", RGBColor(0xC0, 0x39, 0x2B)),
    ("20", "+8.7", "−5.7", "16%", "关注", ORANGE),
    ("30", "+18.7", "+4.3", "1.6%", "安全", GREEN),
    ("80", "+68.7", "+54.3", "约 0", "安全（演示价）", GREEN),
  ]
  xs = [822960, 3100000, 5200000, 7300000, 9400000]
  ws = [2200000, 2000000, 2000000, 2000000, 2000000]
  for x, w, h in zip(xs, ws, heads):
    tb(s, x, 1180000, w, 340000, [(h, 13, True, BLUE)])
  for ri, row in enumerate(data):
    y = 1580000 + ri * 500000
    *cells, jcol = row
    accent = ri == 3
    for x, w, cell in zip(xs, ws, cells):
      last = cell == cells[-1]
      col = jcol if last else (ORANGE if accent else INK)
      tb(s, x, y, w, 460000, [(cell, 16 if accent else 14, accent or last, col)])
  tb(s, 822960, 4200000, 10515600, 1200000, [
    ("变动成本层，装修报价不要低于 30 元/t。20 元期望还能赚，尾部仍可能亏。", 14, True, INK),
    ("演示 80 元时边际 +68.7，不代表 0 费也赚。湿样到位后可用该批入炉湿样替换 720 残差。", 13, False, MUTE),
  ])


def slide_64(prs):
  s = add_blank(prs)
  header(s, 8, "一般固废入炉热值 · 定价模型接入")
  tb(s, 822960, 740000, 10515600, 380000,
     [("敏感性：农林若只有检测一半", 24, True, INK)])
  tb(s, 822960, 1160000, 10515600, 360000, [
    ("2792 → 1396 大卡。锅炉反推的综合 1188 不一定动——化验砍半，改的是热记在谁账上。", 14, False, INK),
  ])
  tb(s, 822960, 1600000, 5120000, 2300000, [
    ("口径 A  篮子总热不变", 16, True, BLUE),
    ("热还给装修", 13, False, MUTE),
    ("综合 0 费边际  仍约 +24", 14, False, INK),
    ("装修残差升至约 1047 大卡", 14, False, INK),
    ("装修 0 费转正约 +15", 18, True, GREEN),
    ("农林仍 +43（高于打平线）", 14, False, INK),
  ], fill=PILL)
  tb(s, 6400000, 1600000, 5120000, 2300000, [
    ("口径 B  总热跟着农林降", 16, True, BLUE),
    ("成分加权", 13, False, MUTE),
    ("综合落到约 987 大卡", 14, False, INK),
    ("综合 0 费约 +7～8", 14, False, INK),
    ("装修仍按 720，0 费仍亏 11", 18, True, ORANGE),
    ("农林仍 +43", 14, False, INK),
  ], fill=PILL)
  tb(s, 822960, 4100000, 10515600, 1400000, [
    ("农林砍一半，它自己 0 费仍赚。真正被农林化验绑架的是装修。", 15, True, INK),
    ("此段是假设，雷博尔报告仍是 2792。要当报价，先定用 A 还是 B。", 13, False, MUTE),
  ])


def slide_interact(prs):
  s = add_blank(prs)
  header(s, 9, "一般固废入炉热值 · 定价模型接入")
  tb(s, 822960, 740000, 10515600, 420000, [("交互测算", 26, True, INK)])
  tb(s, 822960, 1220000, 10515600, 420000, [
    ("打开同文件夹里的  固废定价测算_交互.html", 16, True, BLUE),
  ])
  tb(s, 822960, 1680000, 5120000, 2200000, [
    ("您填", 14, True, BLUE),
    ("处置费、上网电价、炉渣运出价、掺烧比例", 15, False, INK),
    ("表里还可以改热值、含水、灰分、渗滤液", 14, False, INK),
    ("最底下一行是空位，可加其他种类", 14, False, INK),
  ], fill=PILL)
  tb(s, 6400000, 1680000, 5120000, 2200000, [
    ("立刻算出", 14, True, BLUE),
    ("综合、造纸、建装筛上物、农林、其他工业", 15, False, INK),
    ("0 费边际、含费边际、打平价 B", 14, False, INK),
    ("点任一行，下面用大白话拆账", 14, False, INK),
  ], fill=PILL)
  tb(s, 822960, 4100000, 10515600, 1400000, [
    ("给领导当场改数用。变动成本层，不是合同价，未扣折旧人工财务。", 15, True, INK),
    ("综合是整篮混料，不要和四类加总。建装筛上物 = 装修入炉 720 大卡。", 13, False, MUTE),
  ])


def lock(prs):
  s = add_blank(prs, navy=True)
  tb(s, 822960, 900000, 10515600, 500000, [("锁定", 28, True, WHITE)])
  tb(s, 822960, 1500000, 10515600, 900000,
     [("综合 1188 大卡        装修入炉约 720 大卡", 26, True, WHITE)])
  tb(s, 822960, 2800000, 5120000, 1600000, [
    ("综合 · 0 处置费", 14, False, RGBColor(0xA8, 0xB4, 0xC4)),
    ("赚  约 24 元/t", 22, True, RGBColor(0x5C, 0xD6, 0x8A)),
    ("变动成本层。发电盖住药剂。", 13, False, RGBColor(0xA8, 0xB4, 0xC4)),
    ("扣折旧人工后仍亏约 24。", 13, False, RGBColor(0xA8, 0xB4, 0xC4)),
  ])
  tb(s, 6400000, 2800000, 5120000, 1600000, [
    ("装修 · 0 处置费", 14, False, RGBColor(0xA8, 0xB4, 0xC4)),
    ("亏  约 11 元/t", 22, True, ORANGE),
    ("720 大卡不够热，药费按吨走。", 13, False, RGBColor(0xA8, 0xB4, 0xC4)),
    ("报价不宜低于 30 元/t。", 13, False, RGBColor(0xA8, 0xB4, 0xC4)),
  ])
  tb(s, 822960, 4700000, 10515600, 700000, [
    ("两句话不要混：综合 0 费不亏，推不出装修 0 费也不亏。上表是变动成本，不是合同价。", 14, False, RGBColor(0xC4, 0xCD, 0xD6)),
  ])


def main() -> int:
  prs = Presentation()
  prs.slide_width, prs.slide_height = W, H
  cover(prs)
  slide_koujing(prs)
  slide_gongshi(prs)
  slide_zhuangxiu(prs)
  slide_zhuashou(prs)
  slide_61(prs)
  slide_62(prs)
  slide_63(prs)
  slide_64(prs)
  slide_interact(prs)
  lock(prs)
  prs.save(str(OUT))
  print("pptx", OUT, "slides", len(prs.slides))
  return 0


if __name__ == "__main__":
  raise SystemExit(main())
