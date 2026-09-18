# -*- coding: utf-8 -*-
"""入炉反推热值 → 定价表（不覆盖 V4 算例 xlsx）。

装修 720 大卡（用户最新模型）；综合取 result.json 解法 B；
造纸/农林钉检测报告。堆放天数=0（已是入炉态）。
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from pathlib import Path

XLSX_SKILL = ""  # 可选：环境变量 XLSX_SKILL
for p in (XLSX_SKILL, XLSX_SKILL + "/templates"):
    if p not in sys.path:
        sys.path.insert(0, p)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from base import (ACCENT_NEGATIVE, ACCENT_POSITIVE, fill_data_row, fill_header,
                  fill_total, font_header, font_title, setup_sheet)

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from hhv.pricing import (  # noqa: E402
    COST, BatchInput, KJ_PER_KCAL, batch_economics, fixed_cost_per_t,
    load_plant_slag, load_pricing_calibration, portfolio, price_tier_table,
    shutdown_economics,
)

RESULT = ROOT / "outputs" / "config.snmis_daily" / "result.json"
OUT = ROOT / "outputs" / "pricing"
OUT.mkdir(parents=True, exist_ok=True)
XLSX = OUT / "定价表_入炉反推_装修720.xlsx"
PRICES = [0, 10, 20, 30, 40, 50, 60, 70, 80, 90, 100, 120]


def put(ws, row, col, val, fmt=None, header=False, total=False, bold=False):
    cell = ws.cell(row=row, column=col, value=val)
    if header:
        cell.fill = fill_header()
        cell.font = font_header()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    elif total:
        cell.fill = fill_total()
        cell.font = Font(bold=True)
    else:
        cell.fill = fill_data_row(row)
        cell.font = Font(bold=bold)
    if fmt:
        cell.number_format = fmt
    return cell


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = font_title()
    return row + 1


def q_isw_kcal() -> tuple[float, float]:
    res = json.loads(RESULT.read_text(encoding="utf-8"))
    d = res["ci_anchored"]["Q_工业固废"]
    mean = d["mean"] / KJ_PER_KCAL
    sd = (d["hi"] - d["lo"]) / KJ_PER_KCAL / (2 * 1.96)
    return mean, sd


def batches(q_mix: float) -> list[BatchInput]:
    # stock_days=0 → 入炉态不再打堆放折
    return [
        BatchInput("综合固废", q_mix, 30, 30, 25, 0, 20, 80, 0, 0, 0.15, 314,
                   notes="解法B 反推入炉；含水/灰分为组合演示值"),
        BatchInput("装修", 720, 28, 28, 13, 0, 20, 80, 0, 0, 0.08, 192,
                   notes="最新模型入炉 720 大卡；灰分/含水近天越化验，渗滤液沿用建装 0.08"),
        BatchInput("造纸(其它)", 1626, 58, 58, 19, 0, 20, 80, 0, 0, 0.35, 51,
                   notes="理文 2026-02 检测报告收到基"),
        BatchInput("农林", 2792, 33, 33, 12, 0, 20, 60, 0, 0, 0.10, 45,
                   notes="雷博尔 2026-03 检测报告收到基"),
    ]


def main() -> int:
    plant = load_plant_slag()
    if plant is None:
        print("缺炉渣台账 data/plant/slag_outbound.csv")
        return 1
    # 厂内口径统一用实测校准（data/pricing_calibration.yaml），与 PFAiX 定价页同一套数
    cal = load_pricing_calibration()
    if not cal:
        print("⚠ 未找到 data/pricing_calibration.yaml——将按 V4 假设值（eta 78~81%、汽耗率 4.8）计算，"
              "与 PFAiX 定价页口径不一致")
    slag_kw = {"slag_rate_pct": plant["rate_pct"], "calibration": cal}
    slag_note = (f"出厂 {plant['slag_out_t']:.1f} t / 入炉 {plant['furnace_input_t']:.1f} t "
                 f"= {plant['rate_pct']:.2f}%")
    q_mix, sd_mix = q_isw_kcal()
    bs = batches(q_mix)
    days, total_burn = 330.0, 2100.0 + 2300.0 * 0.20
    fixed_pt = fixed_cost_per_t(total_burn, days)
    ecoms = [batch_economics(b, fixed_per_t=fixed_pt, **slag_kw) for b in bs]
    port = portfolio(bs, ecoms)
    shut = shutdown_economics(weighted_margin=port["margin"])

    wb = Workbook()
    ws = wb.active
    ws.title = "01-单批次"
    setup_sheet(ws, title="入炉反推定价（装修 720 大卡 · 堆放折=1）", last_col=6)
    row = 3
    row = section(ws, row, "一、输入")
    heads = ["指标", "单位"] + [b.name for b in bs]
    for j, h in enumerate(heads, 1):
        put(ws, row, j, h, header=True)
    row += 1
    for label, unit, key, fmt in [
        ("入炉低位热值", "kcal/kg", "q_lab", "0"),
        ("含水率（化验）", "%", "moisture_lab", "0"),
        ("灰分", "%", "ash", "0"),
        ("堆放天数", "天", "stock_days", "0"),
        ("掺烧比例", "%", "blend_pct", "0"),
        ("演示处置费", "元/t", "disposal_price", "0.00"),
        ("渗滤液产生率", "吨/吨", "leachate_rate", "0.00"),
        ("日处置量（权重）", "t/d", "daily_t", "0"),
    ]:
        put(ws, row, 1, label)
        put(ws, row, 2, unit)
        for j, b in enumerate(bs, 3):
            put(ws, row, j, getattr(b, key), fmt)
        row += 1
    row += 1
    row = section(ws, row, "二、能量与成本")
    put(ws, row, 1, "炉渣运出价（出厂）")
    put(ws, row, 2, "元/t渣")
    for j in range(3, 3 + len(bs)):
        put(ws, row, j, COST["slag_price"], "0.00")
    row += 1
    for label, unit, key, fmt in [
        ("炉渣率（出厂/入炉）", "%", "slag_rate", "0.00"),
        ("堆放损耗 k", "", "k_stock", "0.00"),
        ("入炉热值（计价）", "kcal/kg", "q_in", "0.0"),
        ("上网发电收入", "元/t垃圾", "power_rev", "0.00"),
        ("炉渣收入（30×渣率）", "元/t垃圾", "slag_rev", "0.00"),
        ("变动成本", "元/t垃圾", "var_cost", "0.00"),
    ]:
        put(ws, row, 1, label)
        put(ws, row, 2, unit)
        for j, e in enumerate(ecoms, 3):
            put(ws, row, j, e[key], fmt)
        row += 1
    row += 1
    row = section(ws, row, "三、决策指标（演示处置费下）")
    for label, unit, key, fmt, tot in [
        ("边际贡献（含处置费）", "元/t", "margin", "0.00", True),
        ("边际（不含处置费）", "元/t", "margin_ex_disposal", "0.00", False),
        ("吨固定成本", "元/t", "fixed_per_t", "0.00", False),
        ("全成本净利", "元/t", "full_cost_profit", "0.00", True),
        ("盈亏平衡处置价 B", "元/t", "be_price", "0.00", True),
    ]:
        put(ws, row, 1, label, total=tot)
        put(ws, row, 2, unit, total=tot)
        for j, e in enumerate(ecoms, 3):
            put(ws, row, j, e[key], fmt, total=tot)
        row += 1

    ws2 = wb.create_sheet("02-组合")
    setup_sheet(ws2, title="四类加权（权重=日处置量）", last_col=3)
    row = 3
    row = section(ws2, row, f"组合 {port['total_daily_t']:.0f} t/d")
    for label, key, unit in [
        ("加权边际", "margin", "元/t"),
        ("加权入炉热值", "q_in", "kcal/kg"),
        ("加权发电收入", "power_rev", "元/t"),
        ("加权处置费", "disposal", "元/t"),
    ]:
        put(ws2, row, 1, label)
        put(ws2, row, 2, port[key], "0.00" if unit != "kcal/kg" else "0.0")
        put(ws2, row, 3, unit)
        row += 1
    row += 1
    put(ws2, row, 1, "协同增量总价值")
    put(ws2, row, 2, shut["synergy_total"], "0.0", total=True)
    put(ws2, row, 3, "万元/年", total=True)

    ws3 = wb.create_sheet("03-报价分档")
    setup_sheet(ws3, title="处置价分档：期望边际 / P5 边际 / 亏损概率", last_col=6)
    row = 3
    sd_map = {
        "综合固废": sd_mix,
        "装修": 720 * 0.15,
        "造纸(其它)": 1626 * 0.08,
        "农林": 2792 * 0.08,
    }
    for b, e in zip(bs, ecoms):
        q_mean = e["q_in"]
        q_sd = max(sd_map[b.name], 1.0)
        tiers = price_tier_table(q_mean, q_sd, b, PRICES, **slag_kw)
        row = section(ws3, row, f"{b.name}｜入炉 {q_mean:.0f} 大卡｜σ={q_sd:.0f}｜盈亏平衡处置价 {e['be_price']:.1f} 元/t")
        for j, h in enumerate(["处置价 元/t", "期望边际 元/t", "P5 边际 元/t", "亏损概率", "判定"], 1):
            put(ws3, row, j, h, header=True)
        row += 1
        for t in tiers:
            put(ws3, row, 1, t["price"], "0")
            put(ws3, row, 2, t["margin_mean"], "0.00")
            put(ws3, row, 3, t["margin_p5"], "0.00")
            put(ws3, row, 4, t["loss_prob"], "0.0%")
            lab = "高风险" if t["loss_prob"] > 0.2 else ("关注" if t["loss_prob"] > 0.05 else "安全")
            cell = put(ws3, row, 5, lab)
            if t["loss_prob"] > 0.2:
                cell.font = Font(color=ACCENT_NEGATIVE, bold=True)
            elif t["loss_prob"] <= 0.05:
                cell.font = Font(color=ACCENT_POSITIVE)
            row += 1
        row += 1

    ws4 = wb.create_sheet("04-口径")
    setup_sheet(ws4, title="口径", last_col=2)
    row = 3
    notes = [
        ("装修热值", "720 kcal/kg，用户最新模型入炉残差；不是化验 4438～5747 大卡。"),
        ("综合固废", f"解法 B {q_mix:.0f} kcal/kg（result.json，约 {q_mix * KJ_PER_KCAL:.0f} kJ/kg）。σ 取 95%CI 反推。"),
        ("造纸 / 农林", "理文底渣 1626 大卡、雷博尔农林 2792 大卡，检测报告收到基。"),
        ("堆放折", "stock_days=0 → k_stock=1。反推已是入炉态，不再打 0.94/0.97。"),
        ("含水", "厂内 k_moist=1。灰分仅影响高灰加成（本表装修/农林/综合灰分均 ≤30%）。"),
        ("炉渣运出价", f"{COST['slag_price']:.0f} 元/t渣（出厂卖价）。吨垃圾收入 = {COST['slag_price']:.0f} × 渣率，不是每吨垃圾进账 {COST['slag_price']:.0f}。"),
        ("炉渣率", slag_note),
        ("演示处置费", "综合/装修/造纸 80 元/t，农林 60 元/t。合同价看 03 分档，不看这一列。"),
        ("B 的含义", "盈亏平衡处置价=变动成本−发电−炉渣。未扣折旧人工财务。B>0 表示没有处置费时变动成本层亏损。"),
        ("生成", f"{datetime.now():%Y-%m-%d %H:%M} · 不覆盖 市场化垃圾定价建议_V4.1.xlsx"),
    ]
    for k, v in notes:
        put(ws4, row, 1, k, bold=True)
        c = put(ws4, row, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws4.row_dimensions[row].height = 28
        row += 1
    ws4.column_dimensions["A"].width = 18
    ws4.column_dimensions["B"].width = 88

    wb.save(XLSX)
    print("written", XLSX)
    print("slag", slag_note)
    print(f"综合 {q_mix:.0f} kcal  σ={sd_mix:.0f}")
    print(f"组合边际 {port['margin']:.1f} 元/t")
    for b, e in zip(bs, ecoms):
        q_sd = max(sd_map[b.name], 1.0)
        tiers = price_tier_table(e["q_in"], q_sd, b, PRICES, **slag_kw)
        risky = [t for t in tiers if t["loss_prob"] > 0.05]
        floor = risky[-1]["price"] if risky else tiers[0]["price"]
        print(f"  {b.name}: q_in={e['q_in']:.0f}  发电{e['power_rev']:.1f}  变动{e['var_cost']:.1f}  "
              f"边际@{b.disposal_price:.0f}={e['margin']:.1f}  B={e['be_price']:.1f}  "
              f"亏>5%下限≈{floor}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
