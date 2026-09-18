# -*- coding: utf-8 -*-
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

raw = Path(__file__).resolve().parents[1] / "data" / "snmis_raw"
for name in ["gdetail_2026-08-01.xlsx", "gdetail_2026-02-15.xlsx"]:
    wb = load_workbook(raw / name, data_only=True)
    print("====", name, wb.sheetnames)
    for sn in wb.sheetnames:
        ws = wb[sn]
        goods = Counter()
        weights = Counter()
        for r in range(3, ws.max_row + 1):
            g = ws.cell(r, 5).value
            w = ws.cell(r, 9).value
            if g is None:
                continue
            gs = str(g).strip()
            if not gs or "合计" in gs:
                continue
            goods[gs] += 1
            try:
                weights[gs] += float(w or 0)
            except (TypeError, ValueError):
                pass
        print("--", sn, "rows", ws.max_row)
        for k, v in goods.most_common(25):
            print(f"  {k}: n={v} t={weights[k]:.1f}")
    print()
